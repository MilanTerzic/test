"""
Read regression coefficients from the source workbook.

Looks at the sheet "Serbian Gas Cons. forecast" and pulls the four polynomial
coefficients and the two linear coefficients from the cells where they live
in the v8 Kalotina workbook.
"""

from __future__ import annotations

from typing import IO, Tuple, Union

import openpyxl

from config import LINEAR_COEFFS, POLY_COEFFS


def read_coefficients_from_xlsx(
    upload: Union[IO, str],
) -> Tuple[Tuple[float, float, float, float], Tuple[float, float]]:
    """Return ((a3, a2, a1, a0), (b1, b0)) — numpy polyval convention."""
    wb = openpyxl.load_workbook(upload, data_only=True)
    if "Serbian Gas Cons. forecast" in wb.sheetnames:
        ws = wb["Serbian Gas Cons. forecast"]

        a0 = ws["T13"].value
        a1 = ws["U13"].value
        a2 = ws["V13"].value
        a3 = ws["W13"].value

        b0 = ws["T14"].value
        b1 = ws["U14"].value

        if None not in (a0, a1, a2, a3):
            poly = (float(a3), float(a2), float(a1), float(a0))
        else:
            poly = POLY_COEFFS
        if None not in (b0, b1):
            lin = (float(b1), float(b0))
        else:
            lin = LINEAR_COEFFS
        return poly, lin

    return POLY_COEFFS, LINEAR_COEFFS
